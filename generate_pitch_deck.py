"""
NiveshRakshak AI — Pitch Deck Generator
Creates a professional, widescreen (16:9) presentation for the SANGYAN Investor Resilience Hackathon.
"""

import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Initialize 16:9 Widescreen Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_slide_layout = prs.slide_layouts[6]

# Theme Colors
BG_COLOR = RGBColor(11, 15, 25)          # Deep Slate/Navy #0B0F19
CARD_BG = RGBColor(22, 30, 49)           # Slate Card #161E31
CARD_BORDER = RGBColor(38, 52, 82)       # Subtle Card Border
PRIMARY_BLUE = RGBColor(37, 99, 235)     # #2563EB
ACCENT_CYAN = RGBColor(56, 189, 248)     # #38BDF8
EMERALD_GREEN = RGBColor(16, 185, 129)   # #10B981
AMBER_WARNING = RGBColor(245, 158, 11)   # #F59E0B
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184)     # #94A3B8
TEXT_LIGHT = RGBColor(226, 232, 240)     # #E2E8F0

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, badge_text, title_text, subtitle_text=""):
    # Category Badge
    badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.4))
    tf_b = badge_box.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = badge_text.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_CYAN

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(26)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE

    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.7), Inches(0.5))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = TEXT_MUTED

def add_card(slide, left, top, width, height, title, items, badge="", accent_color=PRIMARY_BLUE):
    # Background shape
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = CARD_BORDER
    shape.line.width = Pt(1)

    # Content Box
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
    tf = txBox.text_frame
    tf.word_wrap = True

    if badge:
        p_b = tf.paragraphs[0]
        p_b.text = badge.upper()
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = accent_color
        p_t = tf.add_paragraph()
    else:
        p_t = tf.paragraphs[0]

    p_t.text = title
    p_t.font.size = Pt(15)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_WHITE
    p_t.space_after = Pt(10)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide1)

# Hackathon Badge
h_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(0.5))
tf1 = h_box.text_frame
p1 = tf1.paragraphs[0]
p1.text = "SANGYAN INVESTOR RESILIENCE HACKATHON 2026  |  INVESTOR PROTECTION TRACK"
p1.font.size = Pt(12)
p1.font.bold = True
p1.font.color.rgb = ACCENT_CYAN

# Main Title
title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.6))
tf_main = title_box.text_frame
p_m = tf_main.paragraphs[0]
p_m.text = "NiveshRakshak AI (निवेशरक्षक)"
p_m.font.size = Pt(44)
p_m.font.bold = True
p_m.font.color.rgb = TEXT_WHITE

# Subtitle
sub_box = slide1.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.3), Inches(1.2))
tf_sub = sub_box.text_frame
tf_sub.word_wrap = True
p_s = tf_sub.paragraphs[0]
p_s.text = "Explainable Multi-Modal Investor Protection & Financial Misinformation Detection Platform for Indian Retail Investors"
p_s.font.size = Pt(20)
p_s.font.color.rgb = TEXT_MUTED

# Metric Badges on Title
badge_w = Inches(3.5)
badges = [
    ("Vernacular-First", "English, हिन्दी & मराठी Parity", EMERALD_GREEN),
    ("Multi-Channel Shield", "Screenshot, Text, URL & Voice", PRIMARY_BLUE),
    ("Safety Guardrails", "Strictly Non-Advisory / Educational", AMBER_WARNING)
]
for idx, (b_title, b_sub, col) in enumerate(badges):
    bx = Inches(1.0 + idx * 3.8)
    by = Inches(5.2)
    s = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, badge_w, Inches(1.3))
    s.fill.solid()
    s.fill.fore_color.rgb = CARD_BG
    s.line.color.rgb = col
    s.line.width = Pt(1.5)
    tb = slide1.shapes.add_textbox(bx + Inches(0.2), by + Inches(0.2), badge_w - Inches(0.4), Inches(0.9))
    tf = tb.text_frame
    p_bt = tf.paragraphs[0]
    p_bt.text = b_title
    p_bt.font.size = Pt(14)
    p_bt.font.bold = True
    p_bt.font.color.rgb = col
    p_bs = tf.add_paragraph()
    p_bs.text = b_sub
    p_bs.font.size = Pt(11)
    p_bs.font.color.rgb = TEXT_LIGHT

# ==============================================================================
# SLIDE 2: THE PROBLEM STATEMENT
# ==============================================================================
slide2 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide2)
add_header(slide2, "The Core Crisis", "India's Retail Investment Boom vs. Sophisticated Financial Fraud",
           "Over 16 Crore Demat accounts exist today, with 70%+ incremental growth coming from Tier-2/3 cities.")

add_card(slide2, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "Surge in Multi-Channel Scams", [
             "WhatsApp & Telegram 'VIP Jackpot' groups claiming 100% daily guaranteed returns.",
             "Social engineering Demat freeze calls demanding urgent verification OTPs.",
             "Pump-and-dump cartels trapping retail investors in illiquid penny stocks.",
             "Lookalike broker phishing domains mimicking Zerodha, AngelOne, Groww, and SBI."
         ], badge="Threat Landscape", accent_color=AMBER_WARNING)

add_card(slide2, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "The Vulnerability Gap", [
             "First-time investors in Tier-2/3 regions lack formal financial education.",
             "Regulatory disclaimers are written in dense, intimidating English legal jargon.",
             "Finfluencers exploit FOMO using rented luxury cars and inspect-element profit fakes.",
             "Victims frequently lose lifetime savings within 48 hours of initial engagement."
         ], badge="Retail Investor Dilemma", accent_color=PRIMARY_BLUE)

add_card(slide2, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.6),
         "Why Existing Tools Fail", [
             "Generic AI chatbots hallucinate or provide unsolicited buy/sell trading advice.",
             "No unified tool exists to verify WhatsApp screenshots, URLs, and claims in one place.",
             "Senior citizens and rural investors cannot navigate complex desktop portals.",
             "Lack of vernacular speech support (Hindi & Marathi voice queries ignored)."
         ], badge="Current Limitations", accent_color=EMERALD_GREEN)

# ==============================================================================
# SLIDE 3: OUR SOLUTION — NIVESHRAKSHAK AI
# ==============================================================================
slide3 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide3)
add_header(slide3, "The Innovation", "NiveshRakshak AI: A Defense-in-Depth Protection Ecosystem",
           "A dual-engine platform combining deterministic SEBI statutory rules with contextual LLM intelligence.")

add_card(slide3, Inches(0.8), Inches(2.2), Inches(5.6), Inches(2.2),
         "1. Multi-Modal Verification Shield", [
             "Screenshot OCR (Tesseract + PIL) to inspect WhatsApp/Telegram forwards.",
             "URL Safety Engine detecting punycode, broker impersonation & disposable TLDs.",
             "Claim Verifier validating SEBI statutory registration prefixes (INH, INA, INZ, INP)."
         ], badge="Channel Agnostic", accent_color=PRIMARY_BLUE)

add_card(slide3, Inches(6.8), Inches(2.2), Inches(5.7), Inches(2.2),
         "2. Explainable Vernacular Diagnostics", [
             "0-100 Risk Score with explicit red-flag breakdown & point deductions.",
             "Full multilingual parity in English, हिन्दी, and मराठी.",
             "Vernacular Voice Assistant with Web Speech API STT and tailored elder TTS."
         ], badge="Accessible & Transparent", accent_color=EMERALD_GREEN)

add_card(slide3, Inches(0.8), Inches(4.7), Inches(5.6), Inches(2.2),
         "3. Interactive Scam Resilience Simulator", [
             "8 real-world Tier-2/3 scam scenarios testing urgency, FOMO & UPI tricks.",
             "Instant behavioral feedback and shareable Investor Safety Certificate.",
             "Contextual shortcuts to 8 foundational financial education modules."
         ], badge="Proactive Capacity Building", accent_color=AMBER_WARNING)

add_card(slide3, Inches(6.8), Inches(4.7), Inches(5.7), Inches(2.2),
         "4. Regulatory Safety & Ethical Guardrails", [
             "Strictly educational: ZERO stock tips, trading advice, or portfolio management.",
             "Zero Personal Data Stored: No OTPs, passwords, or bank numbers ever persisted.",
             "Emergency 1-click connection to National Cybercrime Helpline (1930)."
         ], badge="Public-Good Architecture", accent_color=ACCENT_CYAN)

# ==============================================================================
# SLIDE 4: CORE VERIFICATION MODULES
# ==============================================================================
slide4 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide4)
add_header(slide4, "Feature Breakdown", "Six Specialized Verification & Protection Engines",
           "Purpose-built micro-engines addressing every touchpoint of retail investor interaction.")

modules = [
    ("A. Screenshot Analyzer", "Tesseract OCR + heuristic scanning extracts text from forwarded images and flags scam patterns instantly.", PRIMARY_BLUE),
    ("B. URL Safety Checker", "Inspects domains for IDN homograph punycode (xn--), broker typosquatting, port exploitation, and HTTP traps.", EMERALD_GREEN),
    ("C. Claim Verifier", "Regex-audits SEBI 12-character registration numbers (INH, INA, INZ, INP) and enforces guaranteed return bans.", ACCENT_CYAN),
    ("D. Content Classifier", "Calculates Promotional vs. Educational scores, flagging VIP funnels, urgency triggers, and penny stock hype.", AMBER_WARNING),
    ("E. Regional Voice AI", "Voice-in, Voice-out assistant with slow speech rate for seniors, delivering vernacular safety verdicts.", PRIMARY_BLUE),
    ("F. Resilience Test", "Practical simulation with 8 Indian fraud scenarios, generating verifiable investor safety badges.", EMERALD_GREEN)
]

for idx, (m_title, m_desc, col) in enumerate(modules):
    row = idx // 3
    col_idx = idx % 3
    x = Inches(0.8 + col_idx * 3.9)
    y = Inches(2.2 + row * 2.5)
    s = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.2))
    s.fill.solid()
    s.fill.fore_color.rgb = CARD_BG
    s.line.color.rgb = col
    s.line.width = Pt(1)

    tb = slide4.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), Inches(3.3), Inches(1.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = m_title
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = m_desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_LIGHT

# ==============================================================================
# SLIDE 5: SYSTEM ARCHITECTURE & TECH STACK
# ==============================================================================
slide5 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide5)
add_header(slide5, "Technical Blueprint", "High-Performance, Zero-Downtime Hybrid Architecture",
           "Engineered with local-first deterministic rule fallbacks to guarantee 100% availability even during network cuts.")

add_card(slide5, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "Frontend (Next.js 16)", [
             "Next.js 16.3.8 App Router with Turbopack for ultra-fast compilation.",
             "React 19 with Tailwind CSS v4 design system and Lucide icons.",
             "Web Speech API for native browser speech-to-text and vernacular synthesis.",
             "19 fully static prerendered routes with zero hydration mismatches."
         ], badge="Client Presentation Tier", accent_color=PRIMARY_BLUE)

add_card(slide5, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "Backend (FastAPI)", [
             "FastAPI (Python 3.13) high-concurrency async REST endpoints.",
             "Rule-Based Statutory Engine + Fallback Heuristics for 100% offline reliability.",
             "Tesseract OCR & PIL image processing pipelines.",
             "SQLite privacy-preserving analytics database (GDPR-ready single & bulk purge)."
         ], badge="Application & Heuristics Tier", accent_color=EMERALD_GREEN)

add_card(slide5, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.6),
         "Safety & Regulatory Layer", [
             "SEBI PFUTP (Prohibition of Fraudulent & Unfair Trade Practices) rules.",
             "SEBI Intermediaries Registry validation rules (INH, INA, INZ, INP).",
             "Strict Content Guardrails: Automatic refusal of trading or buy/sell calls.",
             "Direct integration with National Cybercrime Reporting Portal (1930)."
         ], badge="Governance & Compliance", accent_color=AMBER_WARNING)

# ==============================================================================
# SLIDE 6: EMPIRICAL TESTING & RIGOROUS VALIDATION
# ==============================================================================
slide6 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide6)
add_header(slide6, "Quality & Reliability", "Exhaustive Empirical Validation Suite",
           "Every module backed by automated unit, integration, and end-to-end network tests.")

metrics = [
    ("100%", "Backend Test Suite", "20/20 test cases passing in test_backend.py", EMERALD_GREEN),
    ("100%", "Live Network Suite", "10/10 real HTTP endpoints validated in test_live_system.py", PRIMARY_BLUE),
    ("19 / 19", "Static Routes Built", "Next.js Turbopack compiled with 0 TypeScript errors", ACCENT_CYAN),
    ("3 Languages", "Multilingual Parity", "Zero-loss translation across EN, हिन्दी & मराठी", AMBER_WARNING)
]

for idx, (m_val, m_lbl, m_det, col) in enumerate(metrics):
    x = Inches(0.8 + idx * 2.95)
    y = Inches(2.2)
    s = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.8), Inches(1.8))
    s.fill.solid()
    s.fill.fore_color.rgb = CARD_BG
    s.line.color.rgb = col
    s.line.width = Pt(1.5)

    tb = slide6.shapes.add_textbox(x + Inches(0.15), y + Inches(0.15), Inches(2.5), Inches(1.5))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    p1.text = m_val
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = col
    p2 = tf.add_paragraph()
    p2.text = m_lbl
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p3 = tf.add_paragraph()
    p3.text = m_det
    p3.font.size = Pt(9)
    p3.font.color.rgb = TEXT_MUTED

add_card(slide6, Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.6),
         "Edge Case Resilience", [
             "Handles empty strings, 50,000-character spam dumps, and malformed URLs gracefully.",
             "Punycode homoglyph detection (e.g. xn--zerdha-...) blocks deceptive Cyrillic lookalikes.",
             "Broker typo detection identifies fake subdomains on disposable .click, .xyz, and .top TLDs."
         ], badge="Adversarial Testing", accent_color=PRIMARY_BLUE)

add_card(slide6, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.6),
         "Live Integration Verification", [
             "Validated live against running FastAPI backend on http://127.0.0.1:8000.",
             "Confirmed live Next.js dev server on http://localhost:3000.",
             "Privacy Center deletion endpoint confirmed to purge SQLite tables on demand."
         ], badge="Deployment Readiness", accent_color=EMERALD_GREEN)

# ==============================================================================
# SLIDE 7: USER JOURNEY & WORKFLOW
# ==============================================================================
slide7 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide7)
add_header(slide7, "Operational Flow", "From Suspicious Message to Verified Protection in 5 Seconds",
           "Designed for frictionless adoption by first-time investors without technical hurdles.")

steps = [
    ("Step 1: Receive Message", "User receives a WhatsApp forward promising 500% profit with 'SEBI Approval'.", PRIMARY_BLUE),
    ("Step 2: Submit to Platform", "User uploads screenshot, pastes link, or speaks query into the Voice Assistant.", ACCENT_CYAN),
    ("Step 3: Instant Forensic Audit", "System inspects claims, scans SEBI registration prefixes, and verifies URL integrity.", AMBER_WARNING),
    ("Step 4: Explainable Verdict", "Delivers 0-100 Risk Score, flags red flags, and provides vernacular audio summary.", EMERALD_GREEN),
    ("Step 5: Educate & Shield", "Directs user to Cyber Helpline (1930) or Learn module to prevent future victimhood.", PRIMARY_BLUE)
]

for idx, (s_title, s_desc, col) in enumerate(steps):
    x = Inches(0.8 + idx * 2.38)
    y = Inches(2.5)
    s = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.25), Inches(4.2))
    s.fill.solid()
    s.fill.fore_color.rgb = CARD_BG
    s.line.color.rgb = col
    s.line.width = Pt(1)

    tb = slide7.shapes.add_textbox(x + Inches(0.15), y + Inches(0.2), Inches(1.95), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = f"0{idx+1}"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = col
    p0.space_after = Pt(8)

    p1 = tf.add_paragraph()
    p1.text = s_title
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = s_desc
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_LIGHT

# ==============================================================================
# SLIDE 8: IMPACT, ROADMAP & CONCLUSION
# ==============================================================================
slide8 = prs.slides.add_slide(blank_slide_layout)
set_slide_background(slide8)
add_header(slide8, "Future Vision", "Scalability, Ecosystem Impact & Roadmap",
           "Building India's National Investor Resilience Shield.")

add_card(slide8, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "Social Impact", [
             "Protects retirement savings and household emergency funds from cyber cartels.",
             "Democratizes regulatory transparency for non-English speaking citizens.",
             "Encourages legitimate wealth creation through regulated mutual funds & index SIPs.",
             "Reduces grievance backlogs for SEBI and cyber police stations."
         ], badge="National Public Good", accent_color=EMERALD_GREEN)

add_card(slide8, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.6),
         "Product Roadmap", [
             "WhatsApp & Telegram automated bot integration for 1-click forward verification.",
             "Live automated sync with SEBI SCORES & BSE/NSE API registries.",
             "Browser Extension to dynamically warn against phishing broker URLs in real-time.",
             "Expansion to 10 additional regional languages (Tamil, Telugu, Bengali, Kannada)."
         ], badge="Near-Term Milestones", accent_color=PRIMARY_BLUE)

add_card(slide8, Inches(8.8), Inches(2.2), Inches(3.7), Inches(4.6),
         "Hackathon Submission Summary", [
             "Complete Full-Stack Prototype operational on FastAPI + Next.js 16.",
             "Fully verified across 20+ test scenarios with 100% passing rate.",
             "Available for live demonstration and immediate open-source evaluation.",
             "Developed with passion for the SANGYAN Investor Resilience Hackathon 2026."
         ], badge="Ready for Evaluation", accent_color=AMBER_WARNING)

# Save presentation
output_path = "NiveshRakshak_AI_Presentation.pptx"
prs.save(output_path)
print(f"[SUCCESS] Pitch deck presentation created successfully at: {output_path}")
